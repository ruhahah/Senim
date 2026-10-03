"""Фото, скриншот или файл (PDF, Word, TXT) → текст для проверки (бета).

Для картинок ИИ, который видит изображения (OpenAI или Gemini), только ПЕРЕПИСЫВАЕТ текст.
PDF и Word читаются без ИИ — библиотеками pypdf и python-docx.
Проверка фактов идёт дальше обычным путём — по тексту, который ученик видит и может поправить.
"""
from __future__ import annotations

import base64
import binascii
import re

from .llm import LLMError, get_router

VISION_PROVIDERS = ("openai", "gemini", "openrouter")
ALLOWED_MIME = {"image/jpeg", "image/png", "image/webp"}
MAX_BYTES = 10_000_000          # фото с телефона «как файл» бывают большими
MAX_TEXT = 12000
MAX_PDF_PAGES = 40
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EXT_MIME = {".pdf": "application/pdf", ".docx": DOCX_MIME, ".txt": "text/plain", ".md": "text/plain",
            ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}

OCR_SYSTEM = """Ты распознаёшь текст на фото или скриншоте для системы фактчекинга Senim.
Обычно на изображении ответ чат-бота (ChatGPT, Gemini и др.), страница учебника или текст ученика.
Перепиши текст ДОСЛОВНО, на том же языке (казахский, русский или английский). Ничего не исправляй,
не переводи, не сокращай и не дополняй — даже если в тексте ошибки: их потом проверит Senim.
- Если это переписка с чат-ботом: в "text" — только ответ бота, в "question" — вопрос пользователя, если он виден.
- Пропусти элементы интерфейса: кнопки, меню, время, значки, подписи-ссылки на источники рядом с текстом.
- Абзацы и пункты списка разделяй переносом строки. Ссылки на литературу (авторы, год, DOI) переписывай полностью.
- Если текста нет или его нельзя разобрать — "text": "" и "readable": false.
Верни ТОЛЬКО JSON: {"text": "...", "question": "...", "readable": true}"""


class OCRError(Exception):
    def __init__(self, code: str, message: str = ""):
        super().__init__(message or code)
        self.code = code


_DATA_URL = re.compile(r"^data:(image/[a-z+]+);base64,(.+)$", re.S)
_ANY_DATA_URL = re.compile(r"^data:([\w.+-]+/[\w.+-]+)?(?:;[^,;]*)*;base64,(.+)$", re.S)


def decode_image(data_url: str) -> tuple[str, bytes]:
    """data:image/jpeg;base64,… → (mime, байты). Проверяет тип и размер."""
    m = _DATA_URL.match((data_url or "").strip())
    if not m:
        raise OCRError("bad_image")
    mime = m.group(1).lower()
    if mime not in ALLOWED_MIME:
        raise OCRError("bad_type")
    try:
        raw = base64.b64decode(m.group(2), validate=True)
    except (binascii.Error, ValueError) as e:
        raise OCRError("bad_image") from e
    if len(raw) > MAX_BYTES:
        raise OCRError("too_big")
    if len(raw) < 100:
        raise OCRError("bad_image")
    return mime, raw


async def extract_text(mime: str, raw: bytes) -> dict:
    """Возвращает {"text", "question"}; пустой text — на фото нет читаемого текста."""
    b64 = base64.b64encode(raw).decode("ascii")
    content = [
        {"type": "text", "text": "Распознай текст на изображении."},
        {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
    ]
    try:
        data = await get_router().complete_json(OCR_SYSTEM, content, cache_ns="ocr", only=VISION_PROVIDERS)
    except LLMError as e:
        raise OCRError("llm_failed", str(e)) from e
    if not isinstance(data, dict):
        data = {}
    text = str(data.get("text") or "").strip()[:MAX_TEXT * 2]
    question = str(data.get("question") or "").strip()[:500]
    return {"text": text, "question": question}


# ------------------------------------------------------------------ файлы: PDF, Word, TXT
def decode_file(data_url: str, name: str = "") -> tuple[str, bytes]:
    """data:…;base64,… → (тип, байты). Тип определяем и по расширению: браузеры не всегда его знают."""
    m = _ANY_DATA_URL.match((data_url or "").strip())
    if not m:
        raise OCRError("bad_file")
    mime = (m.group(1) or "").lower()
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in (name or "") else ""
    if mime not in ALLOWED_MIME | {"application/pdf", DOCX_MIME, "text/plain", "text/markdown"}:
        mime = EXT_MIME.get(ext, mime)
    try:
        raw = base64.b64decode(m.group(2), validate=True)
    except (binascii.Error, ValueError) as e:
        raise OCRError("bad_file") from e
    if len(raw) > MAX_BYTES:
        raise OCRError("too_big")
    if not raw:
        raise OCRError("bad_file")
    return mime, raw


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def read_pdf(raw: bytes) -> str:
    import io

    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(raw))
        parts = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            parts.append(page.extract_text() or "")
            if sum(len(p) for p in parts) > MAX_TEXT * 2:
                break
    except Exception as e:  # noqa: BLE001 — битый или зашифрованный PDF
        raise OCRError("bad_file", str(e)) from e
    text = _tidy("\n\n".join(parts))
    if len(text) < 20:
        raise OCRError("scanned_pdf")  # в PDF только картинки — нужен скриншот или фото страницы
    return text


def read_docx(raw: bytes) -> str:
    import io

    import docx

    try:
        d = docx.Document(io.BytesIO(raw))
    except Exception as e:  # noqa: BLE001
        raise OCRError("bad_file", str(e)) from e
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return _tidy("\n".join(parts))


def read_txt(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return _tidy(raw.decode(enc))
        except UnicodeDecodeError:
            continue
    return _tidy(raw.decode("utf-8", errors="replace"))


async def extract_any(mime: str, raw: bytes) -> dict:
    """Любой поддерживаемый файл → {"text", "question", "kind", "truncated"}."""
    if mime in ALLOWED_MIME:
        out = await extract_text(mime, raw)
        kind = "image"
    else:
        if mime == "application/pdf":
            text, kind = read_pdf(raw), "pdf"
        elif mime == DOCX_MIME:
            text, kind = read_docx(raw), "docx"
        elif mime in ("text/plain", "text/markdown"):
            text, kind = read_txt(raw), "txt"
        else:
            raise OCRError("bad_type")
        out = {"text": text, "question": ""}
    truncated = len(out["text"]) > MAX_TEXT
    out["text"] = out["text"][:MAX_TEXT]
    return {**out, "kind": kind, "truncated": truncated}
